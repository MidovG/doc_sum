"""Извлечение текста из файлов разных форматов."""

from __future__ import annotations

import csv
import html
import io
import json
import logging
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Iterator

log = logging.getLogger(__name__)


@dataclass
class Document:
    path: str            # путь относительно корня папки
    kind: str            # расширение / тип
    text: str = ""
    error: str | None = None
    meta: dict = field(default_factory=dict)

    @property
    def ok(self) -> bool:
        return self.error is None and bool(self.text.strip())


# --------------------------------------------------------------------------- #
# Экстракторы. Каждый принимает bytes и возвращает текст.
# --------------------------------------------------------------------------- #

def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1251", "koi8-r", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _txt(data: bytes) -> str:
    return _decode(data)


def _json(data: bytes) -> str:
    try:
        return json.dumps(json.loads(_decode(data)), ensure_ascii=False, indent=1)
    except json.JSONDecodeError:
        return _decode(data)


def _html(data: bytes) -> str:
    s = _decode(data)
    s = re.sub(r"(?is)<(script|style).*?</\1>", " ", s)
    s = re.sub(r"(?i)<br\s*/?>|</p>|</div>|</h\d>|</li>|</tr>", "\n", s)
    s = re.sub(r"<[^>]+>", " ", s)
    return html.unescape(s)


def _csv(data: bytes) -> str:
    text = _decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096])
    except csv.Error:
        dialect = csv.excel
    rows = csv.reader(io.StringIO(text), dialect)
    return "\n".join(" | ".join(r) for r in rows)


def _pdf(data: bytes) -> str:
    import pymupdf as fitz

    parts = []
    with fitz.open(stream=data, filetype="pdf") as doc:
        for i, page in enumerate(doc, 1):
            t = page.get_text("text").strip()
            if t:
                parts.append(f"[стр. {i}]\n{t}")
    if not parts:
        raise ValueError("в PDF нет текстового слоя (скан) — нужен OCR")
    return "\n\n".join(parts)


def _docx(data: bytes) -> str:
    import docx

    d = docx.Document(io.BytesIO(data))
    out: list[str] = []
    # Абзацы и таблицы в порядке следования в документе
    body = d.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            t = "".join(n.text or "" for n in child.iter() if n.tag.endswith("}t"))
            if t.strip():
                out.append(t)
        elif tag == "tbl":
            for tr in child.iter():
                if tr.tag.endswith("}tr"):
                    cells = []
                    for tc in tr.iter():
                        if tc.tag.endswith("}tc"):
                            cells.append(
                                "".join(n.text or "" for n in tc.iter() if n.tag.endswith("}t")).strip()
                            )
                    out.append(" | ".join(cells))
    return "\n".join(out)


def _xlsx(data: bytes) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"## Лист: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            vals = ["" if v is None else str(v) for v in row]
            if any(v.strip() for v in vals):
                out.append(" | ".join(vals).rstrip(" |"))
    wb.close()
    return "\n".join(out)


def _pptx(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    out = []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
            elif getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    texts.append(" | ".join(c.text for c in row.cells))
        if slide.has_notes_slide:
            n = slide.notes_slide.notes_text_frame.text.strip()
            if n:
                texts.append(f"(заметки: {n})")
        if any(t.strip() for t in texts):
            out.append(f"[слайд {i}]\n" + "\n".join(t for t in texts if t.strip()))
    return "\n\n".join(out)


def _rtf(data: bytes) -> str:
    from striprtf.striprtf import rtf_to_text

    return rtf_to_text(_decode(data))


def _eml(data: bytes) -> str:
    import email
    from email import policy

    msg = email.message_from_bytes(data, policy=policy.default)
    head = "\n".join(f"{h}: {msg[h]}" for h in ("From", "To", "Date", "Subject") if msg[h])
    body = msg.get_body(preferencelist=("plain", "html"))
    content = body.get_content() if body else ""
    if body is not None and body.get_content_type() == "text/html":
        content = _html(content.encode())
    return f"{head}\n\n{content}"


EXTRACTORS: dict[str, Callable[[bytes], str]] = {
    **dict.fromkeys(["txt", "md", "markdown", "log", "rst", "xml", "yaml", "yml", "ini"], _txt),
    "json": _json,
    **dict.fromkeys(["html", "htm"], _html),
    **dict.fromkeys(["csv", "tsv"], _csv),
    "pdf": _pdf,
    "docx": _docx,
    **dict.fromkeys(["xlsx", "xlsm"], _xlsx),
    "pptx": _pptx,
    "rtf": _rtf,
    "eml": _eml,
}

UNSUPPORTED_HINT = {
    "doc": "старый формат Word — сохраните как .docx",
    "xls": "старый формат Excel — сохраните как .xlsx",
    "ppt": "старый формат PowerPoint — сохраните как .pptx",
    **dict.fromkeys(["png", "jpg", "jpeg", "gif", "bmp", "tiff", "webp"], "изображение — нужен OCR"),
}

MAX_ZIP_DEPTH = 2


def extract_bytes(name: str, data: bytes, depth: int = 0) -> Iterator[Document]:
    ext = Path(name).suffix.lower().lstrip(".")
    if ext == "zip":
        if depth >= MAX_ZIP_DEPTH:
            yield Document(name, ext, error="слишком глубокая вложенность архивов")
            return
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                for info in zf.infolist():
                    if info.is_dir() or _is_junk(info.filename):
                        continue
                    yield from extract_bytes(f"{name}/{info.filename}", zf.read(info), depth + 1)
        except zipfile.BadZipFile as e:
            yield Document(name, ext, error=f"битый архив: {e}")
        return

    fn = EXTRACTORS.get(ext)
    if fn is None:
        yield Document(name, ext or "?", error=UNSUPPORTED_HINT.get(ext, "формат не поддерживается"))
        return
    try:
        text = _normalize(fn(data))
        doc = Document(name, ext, text=text, meta={"chars": len(text)})
        if not text.strip():
            doc.error = "текст не найден"
        yield doc
    except Exception as e:  # noqa: BLE001 — один плохой файл не должен ронять весь прогон
        log.warning("Не удалось прочитать %s: %s", name, e)
        yield Document(name, ext, error=str(e) or type(e).__name__)


def _is_junk(name: str) -> bool:
    base = Path(name).name
    return base.startswith(("~$", ".")) or "__MACOSX" in name or base in {"Thumbs.db", "desktop.ini"}


def _normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\xa0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n\s*\n+", "\n\n", text)
    return text.strip()


def load_folder(root: Path) -> list[Document]:
    """Рекурсивно обходит папку (или один файл) и извлекает текст из всех файлов."""
    root = Path(root)
    files = [root] if root.is_file() else sorted(p for p in root.rglob("*") if p.is_file())
    docs: list[Document] = []
    for p in files:
        rel = p.name if root.is_file() else str(p.relative_to(root))
        if _is_junk(rel):
            continue
        docs.extend(extract_bytes(rel, p.read_bytes()))
    return docs
