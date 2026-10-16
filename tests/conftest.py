import io
import zipfile
from pathlib import Path

import pytest

@pytest.fixture()
def sample_folder(tmp_path: Path) -> Path:
    import docx
    import pymupdf as fitz
    import openpyxl
    from pptx import Presentation

    root = tmp_path / "docs"
    (root / "sub").mkdir(parents=True)

    (root / "readme.txt").write_text("Проект «Альфа». Срок сдачи — 1 декабря 2026 года.", encoding="utf-8")
    (root / "cp1251.txt").write_bytes("Договор поставки № 17 от 03.03.2026.".encode("cp1251"))

    d = docx.Document()
    d.add_paragraph("Договор оказания услуг между ООО «Ромашка» и ИП Иванов.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Сумма", "500 000 руб."
    t.cell(1, 0).text, t.cell(1, 1).text = "Срок", "30 дней"
    d.save(root / "contract.docx")

    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Invoice 42. Total due: 1000 USD.")
    pdf.save(root / "sub" / "invoice.pdf")

    wb = openpyxl.Workbook()
    wb.active.append(["Месяц", "Выручка"])
    wb.active.append(["Январь", 120])
    wb.save(root / "sub" / "report.xlsx")

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Roadmap"
    s.placeholders[1].text = "Q1: MVP. Q2: launch."
    prs.save(root / "deck.pptx")

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("inner/notes.md", "# Заметки\nВстреча с заказчиком перенесена.")
        zf.writestr("__MACOSX/._junk", "x")
    (root / "archive.zip").write_bytes(buf.getvalue())

    make_text_image("INVOICE 42 TOTAL 1000").save(root / "photo.png")
    (root / "broken.jpg").write_bytes(b"\xff\xd8 not really a jpeg")

    scan = fitz.open()                          # PDF без текстового слоя — «скан»
    page = scan.new_page()
    buf = io.BytesIO()
    make_text_image("SCANNED SERVICE CONTRACT", size=(1100, 200)).save(buf, "PNG")
    page.insert_image(page.rect, stream=buf.getvalue())
    scan.save(root / "sub" / "scan.pdf")
    (root / "~$lock.docx").write_bytes(b"junk")
    return root


def make_text_image(text: str, size=(900, 200)):
    """Белая картинка с крупным чёрным текстом — пригодна и для vision, и для OCR."""
    from PIL import Image, ImageDraw, ImageFont

    im = Image.new("RGB", size, "white")
    ImageDraw.Draw(im).text((30, 60), text, fill="black", font=ImageFont.load_default(size=56))
    return im
